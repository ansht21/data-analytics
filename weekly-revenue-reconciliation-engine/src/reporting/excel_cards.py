"""Dark "Command Center" card renderers for the openpyxl dashboard.

Provides the Excel equivalents of the Streamlit cockpit components:
gradient hero cards, compact stat pills, section headers and the audit
status banner.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.worksheet import Worksheet

from .excel_theme import (
    ACCENT,
    BANNER_FILL,
    BANNER_TITLE_FONT,
    CENTER,
    GRAD_HERO_A,
    GRAD_HERO_B,
    HERO_FOOT_FONT,
    HERO_LABEL_FONT,
    HERO_SUB_FONT,
    HERO_VALUE_FONT,
    LEFT,
    MEDIUM,
    PILL_FILL,
    PILL_LABEL_FONT,
    PILL_VALUE_FONT,
    SECTION_FONT,
    TEXT_LIGHT,
    hero_gradient,
    section_fill,
)


def safe_merge(
    ws: Worksheet,
    min_row: int,
    min_col: int,
    max_row: int,
    max_col: int,
) -> None:
    """Merge a cell range, refusing to overlap an existing merge.

    Excel treats overlapping merges as file corruption and then refuses to
    open the workbook, so the clash is raised here instead of silently
    producing a file the user cannot open.

    Args:
        ws: Target worksheet.
        min_row: First row of the range.
        min_col: First column of the range.
        max_row: Last row of the range.
        max_col: Last column of the range.
    """
    for merged in list(ws.merged_cells.ranges):
        overlaps = not (
            merged.max_row < min_row or merged.min_row > max_row
            or merged.max_col < min_col or merged.min_col > max_col
        )
        if overlaps:
            raise ValueError(
                f"Merge overlap: {min_col},{min_row}-{max_col},{max_row} "
                f"collides with existing merge {merged.coord}"
            )
    ws.merge_cells(start_row=min_row, start_column=min_col,
                   end_row=max_row, end_column=max_col)


def render_hero_card(
    ws: Worksheet,
    start_row: int,
    start_col: int,
    span: int,
    height: int,
    label: str,
    value: float | str,
    sub: str,
    foot: str = "",
    alt: bool = False,
    number_format: str | None = None,
) -> None:
    """Render a large gradient hero card.

    The card occupies ``height`` rows and ``span`` columns. Every cell in
    the block gets the gradient fill so the card reads as one surface, with
    a bright accent bar on top.

    Args:
        ws: Target worksheet.
        start_row: Top row of the card.
        start_col: Left column of the card.
        span: Number of columns the card covers.
        height: Number of rows the card covers.
        label: Small uppercase label.
        value: Primary value (numeric cells get a number format applied).
        sub: Descriptor under the value.
        foot: Optional footnote line.
        alt: Use the emerald gradient instead of the cyan one.
        number_format: Excel number format for a numeric value.
    """
    stops = GRAD_HERO_B if alt else GRAD_HERO_A
    fill = hero_gradient(stops)
    end_row = start_row + height - 1
    end_col = start_col + span - 1

    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill
            cell.border = Border(
                top=MEDIUM if r == start_row else Side(style=None),
                bottom=Side(style=None),
                left=Side(style=None),
                right=Side(style=None),
            )

    safe_merge(ws, start_row, start_col, start_row, end_col)
    lbl = ws.cell(row=start_row, column=start_col, value=label.upper())
    lbl.font = HERO_LABEL_FONT
    lbl.alignment = LEFT

    safe_merge(ws, start_row + 1, start_col, start_row + 2, end_col)
    val = ws.cell(row=start_row + 1, column=start_col, value=value)
    val.font = HERO_VALUE_FONT
    val.alignment = Alignment(horizontal="left", vertical="center")
    if isinstance(value, (int, float)) and number_format:
        val.number_format = number_format

    safe_merge(ws, start_row + 3, start_col, start_row + 3, end_col)
    subc = ws.cell(row=start_row + 3, column=start_col, value=sub)
    subc.font = HERO_SUB_FONT
    subc.alignment = LEFT

    if foot:
        safe_merge(ws, start_row + 4, start_col, start_row + 4, end_col)
        footc = ws.cell(row=start_row + 4, column=start_col, value=foot)
        footc.font = HERO_FOOT_FONT
        footc.alignment = LEFT


def render_stat_pill(
    ws: Worksheet,
    row: int,
    start_col: int,
    span: int,
    label: str,
    value: str,
    tone: str = "default",
) -> None:
    """Render one compact stat pill (label above a value chip).

    Args:
        ws: Target worksheet.
        row: Top row of the pill.
        start_col: Left column of the pill.
        span: Number of columns the pill covers.
        label: Pill label.
        value: Value text.
        tone: ``pos``, ``warn``, ``neg`` or ``default`` for the value colour.
    """
    colors = {
        "pos": "34D399",
        "warn": "F59E0B",
        "neg": "FB7185",
        "default": TEXT_LIGHT,
    }
    end_col = start_col + span - 1
    for r in (row, row + 1):
        for c in range(start_col, end_col + 1):
            ws.cell(row=r, column=c).fill = PILL_FILL

    safe_merge(ws, row, start_col, row, end_col)
    lbl = ws.cell(row=row, column=start_col, value=label.upper())
    lbl.font = PILL_LABEL_FONT
    lbl.alignment = LEFT

    safe_merge(ws, row + 1, start_col, row + 1, end_col)
    val = ws.cell(row=row + 1, column=start_col, value=value)
    val.font = Font(
        name=PILL_VALUE_FONT.name, size=PILL_VALUE_FONT.size, bold=True,
        color=colors.get(tone, TEXT_LIGHT),
    )
    val.alignment = LEFT


def render_section_header(
    ws: Worksheet,
    row: int,
    start_col: int,
    span: int,
    title: str,
) -> None:
    """Render a section header row.

    Args:
        ws: Target worksheet.
        row: Header row.
        start_col: Left column.
        span: Number of columns the header covers.
        title: Section title.
    """
    end_col = start_col + span - 1
    for c in range(start_col, end_col + 1):
        ws.cell(row=row, column=c).fill = section_fill()
    safe_merge(ws, row, start_col, row, end_col)
    cell = ws.cell(row=row, column=start_col, value=title.upper())
    cell.font = SECTION_FONT
    cell.alignment = LEFT


def render_status_banner(
    ws: Worksheet,
    row: int,
    start_col: int,
    span: int,
    title: str,
    sub: str,
) -> None:
    """Render the audit-status banner.

    Args:
        ws: Target worksheet.
        row: Banner row.
        start_col: Left column.
        span: Number of columns the banner covers.
        title: Bold banner title.
        sub: Supporting line.
    """
    end_col = start_col + span - 1
    for c in range(start_col, end_col + 1):
        ws.cell(row=row, column=c).fill = BANNER_FILL
        ws.cell(row=row, column=c).border = Border(
            top=Side(style="thin", color=ACCENT),
            bottom=Side(style="thin", color=ACCENT),
            left=Side(style="thin", color=ACCENT) if c == start_col else Side(style=None),
            right=Side(style="thin", color=ACCENT) if c == end_col else Side(style=None),
        )
    safe_merge(ws, row, start_col, row, end_col)
    cell = ws.cell(row=row, column=start_col, value=f"{title}   |   {sub}")
    cell.font = BANNER_TITLE_FONT
    cell.alignment = CENTER
