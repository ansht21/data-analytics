"""Exception Audit Log tab: cancelled / returned / lost order detail."""

from __future__ import annotations

import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook

from .formatters import (
    CENTER,
    INR_FORMAT,
    LEFT,
    SUBTITLE_FONT,
    TITLE_FONT,
    style_header_row,
    style_range_border,
)

AUDIT_COLS = [
    "Order ID", "parsed_date", "week_label", "Status",
    "financial_bucket", "Amount", "Fulfilment", "ship-state", "Category",
]


def build_audit_tab(wb: Workbook, exceptions: pd.DataFrame, run_str: str) -> None:
    """Create the pre-filtered Exception Audit Log tab.

    Args:
        wb: Target workbook.
        exceptions: Cancelled / returned / lost-damaged order rows.
        run_str: Human-readable run timestamp for the subtitle.
    """
    audit = wb.create_sheet("Exception Audit Log")
    audit["A1"] = "Exception Audit Log — Cancelled / Returned / Lost-Damaged Orders"
    audit["A1"].font = TITLE_FONT
    audit["A2"] = f"Run Date: {run_str}   |   Rows: {len(exceptions)}"
    audit["A2"].font = SUBTITLE_FONT
    cols = [c for c in AUDIT_COLS if c in exceptions.columns]
    for c, header in enumerate(cols, start=1):
        audit.cell(row=4, column=c, value=header)
    style_header_row(audit, 4, len(cols))
    for i, (_, r) in enumerate(exceptions.iterrows()):
        row = 5 + i
        for c, col in enumerate(cols, start=1):
            val = r[col]
            cell = audit.cell(
                row=row, column=c,
                value=(str(val) if not isinstance(val, (int, float)) else val))
            cell.alignment = LEFT if c <= 5 else CENTER
            if col == "Amount":
                cell.number_format = INR_FORMAT
    last_row = 4 + len(exceptions)
    audit.conditional_formatting.add(
        f"D5:D{max(last_row, 5)}",
        CellIsRule(operator="equal", formula=['"Cancelled"'],
                   fill=PatternFill("solid", fgColor="FFC7CE")))
    style_range_border(audit, 4, max(last_row, 4), 1, len(cols))
    for c in range(1, len(cols) + 1):
        audit.column_dimensions[get_column_letter(c)].width = 22
    audit.freeze_panes = "A5"
    if len(exceptions):
        audit.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{last_row}"
