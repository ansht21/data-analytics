"""Modern 2026 executive KPI card renderer for openpyxl dashboards.

Replaces saturated block colors with clean floating white cards:
- Colored top-accent border
- Muted slate category label
- High-contrast bold metric value in Segoe UI
- Subtle secondary context note
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Side
from openpyxl.worksheet.worksheet import Worksheet

from .formatters import (
    CARD_BG_FILL,
    CARD_BORDER_COLOR,
    CARD_LABEL_FONT,
    CARD_SUB_FONT,
    CARD_VALUE_FONT,
    KPI_ACCENTS,
)


def render_kpi_card(
    ws: Worksheet,
    start_row: int,
    start_col: int,
    label: str,
    value: float | int | str,
    color: str,
    number_format: str = '#,##0.00',
    subtext: str | None = None,
) -> None:
    """Render a modern executive card occupying 3 rows x 2 columns.

    Args:
        ws: Target worksheet.
        start_row: Top row of the card.
        start_col: Left column of the card.
        label: Card category label (row 1).
        value: Primary metric number (row 2).
        color: Accent key into :data:`KPI_ACCENTS` (blue/green/red/yellow/gray/purple).
        number_format: Excel number format string.
        subtext: Context note below the value (row 3).
    """
    accent_hex = KPI_ACCENTS.get(color, "0B2A4A")
    top_side = Side(style="medium", color=accent_hex)
    thin_side = Side(style="thin", color=CARD_BORDER_COLOR)
    hairline = Side(style="hair", color="E2E8F0")

    # Merge cells across 2 columns for each row
    for r_offset in range(3):
        ws.merge_cells(
            start_row=start_row + r_offset,
            start_column=start_col,
            end_row=start_row + r_offset,
            end_column=start_col + 1,
        )

    # 1. Label Row (start_row)
    lbl_cell = ws.cell(row=start_row, column=start_col, value=label.upper())
    lbl_cell.font = CARD_LABEL_FONT
    lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
    lbl_cell.fill = CARD_BG_FILL

    # 2. Value Row (start_row + 1)
    val_cell = ws.cell(row=start_row + 1, column=start_col, value=value)
    val_cell.font = CARD_VALUE_FONT
    val_cell.alignment = Alignment(horizontal="center", vertical="center")
    val_cell.fill = CARD_BG_FILL
    if isinstance(value, (int, float)):
        val_cell.number_format = number_format

    # 3. Subtext Row (start_row + 2)
    sub_cell = ws.cell(row=start_row + 2, column=start_col, value=subtext or "")
    sub_cell.font = CARD_SUB_FONT
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    sub_cell.fill = CARD_BG_FILL

    # Apply precise borders around the 3x2 card perimeter
    for r in range(start_row, start_row + 3):
        for c in range(start_col, start_col + 2):
            cell = ws.cell(row=r, column=c)
            cell.fill = CARD_BG_FILL
            # Calculate outer borders
            t = top_side if r == start_row else (hairline if r == start_row + 1 else Side(style=None))
            b = thin_side if r == start_row + 2 else Side(style=None)
            l_side = thin_side if c == start_col else Side(style=None)
            r_side = thin_side if c == start_col + 1 else Side(style=None)
            cell.border = Border(top=t, bottom=b, left=l_side, right=r_side)
