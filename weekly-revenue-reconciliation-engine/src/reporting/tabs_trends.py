"""Weekly Trends tab: live Excel formulas, conditional formats, data bars."""

from __future__ import annotations

import pandas as pd
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .formatters import (
    CENTER,
    INR_FORMAT,
    INT_FORMAT,
    style_header_row,
    style_range_border,
)

TREND_HEADERS = [
    "Week", "Gross Revenue", "Net Realized", "Cancellation Leakage",
    "Return Leakage", "Loss Leakage", "Total Leakage", "Adjusted Net Cash",
    "Clearance %", "WoW Net %", "Orders", "Cancelled Orders",
    "Cancellation %", "AOV", "Settled Base", "Settled Clearance %",
    "Days Observed", "Partial Week",
]

# Column index of "Clearance %" — the Excel combo chart plots this on axis 2.
CLEARANCE_COL = 9


def build_trends_tab(wb: Workbook, weekly: pd.DataFrame) -> tuple[Worksheet, int, int]:
    """Create the Weekly Trends matrix with live Excel formulas.

    Args:
        wb: Target workbook.
        weekly: Weekly KPI frame from the reconciliation engine.

    Returns:
        Tuple of (sheet, first_data_row, last_data_row).
    """
    trends = wb.create_sheet("Weekly Trends")
    for c, header in enumerate(TREND_HEADERS, start=1):
        trends.cell(row=3, column=c, value=header)
    style_header_row(trends, 3, len(TREND_HEADERS))
    first_row = 4
    for i, (_, r) in enumerate(weekly.iterrows()):
        row = first_row + i
        trends.cell(row=row, column=1, value=str(r["week_label"]))
        trends.cell(row=row, column=2, value=float(r["gross_revenue"]))
        trends.cell(row=row, column=3, value=float(r["net_realized"]))
        trends.cell(row=row, column=4, value=float(r["cancellation_leakage"]))
        trends.cell(row=row, column=5, value=float(r["return_leakage"]))
        trends.cell(row=row, column=6, value=float(r["loss_leakage"]))
        trends.cell(row=row, column=7, value=float(r["total_leakage"]))
        trends.cell(row=row, column=8, value=f"=IFERROR(B{row}-G{row},0)")
        trends.cell(row=row, column=9, value=f"=IFERROR(C{row}/B{row},0)")
        trends.cell(row=row, column=10,
                    value=f"=IFERROR((C{row}-C{row - 1})/C{row - 1},0)" if i else 0.0)
        trends.cell(row=row, column=11, value=int(r["total_orders"]))
        trends.cell(row=row, column=12, value=int(r["cancelled_orders"]))
        trends.cell(row=row, column=13, value=f"=IFERROR(L{row}/K{row},0)")
        trends.cell(row=row, column=14, value=f"=IFERROR(B{row}/K{row},0)")
        # Settled base = Net Realized + Total Leakage (excludes In-Transit),
        # so settled clearance measures real leakage, not settlement lag.
        trends.cell(row=row, column=15, value=f"=C{row}+G{row}")
        trends.cell(row=row, column=16, value=f"=IFERROR(C{row}/O{row},0)")
        trends.cell(row=row, column=17, value=int(r.get("days_observed", 7)))
        trends.cell(row=row, column=18,
                    value="YES" if bool(r.get("is_partial_week", False)) else "NO")
    last_row = first_row + len(weekly) - 1
    total_row = last_row + 1
    trends.cell(row=total_row, column=1, value="TOTAL")
    for col in (2, 3, 4, 5, 6, 7, 8, 11, 12, 15):
        letter = get_column_letter(col)
        trends.cell(row=total_row, column=col,
                    value=f"=SUM({letter}{first_row}:{letter}{last_row})")
    trends.cell(row=total_row, column=9, value=f"=IFERROR(C{total_row}/B{total_row},0)")
    trends.cell(row=total_row, column=10, value="-")
    trends.cell(row=total_row, column=13, value=f"=IFERROR(L{total_row}/K{total_row},0)")
    trends.cell(row=total_row, column=14, value=f"=IFERROR(B{total_row}/K{total_row},0)")
    trends.cell(row=total_row, column=16, value=f"=IFERROR(C{total_row}/O{total_row},0)")
    trends.cell(row=total_row, column=17, value="-")
    trends.cell(row=total_row, column=18, value="-")
    for row in range(first_row, total_row + 1):
        for col, fmt in ((2, INR_FORMAT), (3, INR_FORMAT), (4, INR_FORMAT),
                         (5, INR_FORMAT), (6, INR_FORMAT), (7, INR_FORMAT),
                         (8, INR_FORMAT), (9, "0.00%"), (10, "0.00%"),
                         (11, INT_FORMAT), (12, INT_FORMAT),
                         (13, "0.00%"), (14, INR_FORMAT), (15, INR_FORMAT),
                         (16, "0.00%"), (17, INT_FORMAT)):
            cell = trends.cell(row=row, column=col)
            if not (row == total_row and col == 10):
                cell.number_format = fmt
            cell.alignment = CENTER
        trends.cell(row=row, column=1).alignment = CENTER
        trends.cell(row=row, column=18).alignment = CENTER
    for c in range(1, len(TREND_HEADERS) + 1):
        cell = trends.cell(row=total_row, column=c)
        cell.fill = PatternFill("solid", fgColor="0B2A4A")
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    style_range_border(trends, 3, total_row, 1, len(TREND_HEADERS))
    # Conditional formatting targets the audit metric (settled clearance, col P).
    trends.conditional_formatting.add(
        f"P{first_row}:P{last_row}",
        CellIsRule(operator="lessThan", formula=["0.70"],
                   fill=PatternFill("solid", fgColor="E2E8F0")))
    trends.conditional_formatting.add(
        f"P{first_row}:P{last_row}",
        CellIsRule(operator="greaterThanOrEqual", formula=["0.85"],
                   fill=PatternFill("solid", fgColor="D5E5F5")))
    trends.conditional_formatting.add(
        f"B{first_row}:B{last_row}",
        DataBarRule(start_type="num", start_value=0, end_type="max", color="1F4E79"))
    for c in range(1, len(TREND_HEADERS) + 1):
        trends.column_dimensions[get_column_letter(c)].width = 19
    trends.column_dimensions["A"].width = 12
    trends.freeze_panes = "A4"
    trends.auto_filter.ref = f"A3:{get_column_letter(len(TREND_HEADERS))}{total_row}"
    return trends, first_row, last_row
