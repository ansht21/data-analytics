"""Build ONE all-in-one Excel file from the full Amazon dataset.

Tabs: README | All Data (cleaned) | Weekly Summary | Leakage | Exceptions.
Run:  python make_excel.py
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw" / "Amazon_Sale_Report.csv"
OUT = ROOT / "reports" / "Amazon_All_Data_SINGLE_FILE.xlsx"

BUCKETS = {
    "Net Realized": ["Shipped - Delivered to Buyer"],
    "Cancellation Leakage": ["Cancelled"],
    "Return Leakage": ["Shipped - Returned to Seller", "Shipped - Returning to Seller",
                       "Shipped - Rejected by Buyer"],
    "Loss Leakage": ["Shipped - Lost in Transit", "Shipped - Damaged"],
    "In-Transit": ["Shipped", "Shipped - Picked Up", "Shipped - Out for Delivery", "Shipping"],
    "Pending": ["Pending", "Pending - Waiting for Pick Up"],
}
LOOKUP = {s: b for b, lst in BUCKETS.items() for s in lst}

HDR_FILL = PatternFill("solid", fgColor="0D2137")
HDR_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
TITLE = Font(name="Calibri", size=16, bold=True, color="0D2137")
CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_header(ws, row: int, ncols: int) -> None:
    """Paint one dark header row."""
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CTR


def clean(df: pd.DataFrame) -> pd.DataFrame:
    """Apply the 11-step compact cleaning + fiscal + bucket mapping."""
    df = df.drop(columns=[c for c in ["index", "New", "PendingS"] if c in df.columns])
    df["parsed_date"] = pd.to_datetime(df["Date"], format="mixed", dayfirst=False)
    df["Amount"] = df["Amount"].fillna(0.0).astype(float)
    if "currency" in df.columns:
        df["currency"] = df["currency"].fillna("INR")
    for col in ("ship-state", "ship-city"):
        if col in df.columns:
            df[col] = df[col].fillna("UNKNOWN").astype(str).str.strip().str.upper()
    if "ship-postal-code" in df.columns:
        df["ship-postal-code"] = pd.to_numeric(
            df["ship-postal-code"], errors="coerce").fillna(0).astype(int)
    if "fulfilled-by" in df.columns:
        df.loc[(df["Fulfilment"] == "Amazon") & df["fulfilled-by"].isna(),
               "fulfilled-by"] = "Amazon FBA"
    iso = df["parsed_date"].dt.isocalendar()
    df["fiscal_week"] = iso["week"].astype(int)
    df["fiscal_year"] = iso["year"].astype(int)
    df["week_label"] = "W" + df["fiscal_week"].astype(str)
    df["financial_bucket"] = df["Status"].astype(str).map(lambda s: LOOKUP.get(s, "Unmapped"))
    return df


def weekly_summary(df: pd.DataFrame) -> pd.DataFrame:
    """Aggregate gross/net/leakage KPIs per fiscal week."""
    rows = []
    for (lbl, _yr, _wk), g in df.groupby(["week_label", "fiscal_year", "fiscal_week"]):
        gross = float(g["Amount"].sum())
        net = float(g.loc[g["financial_bucket"] == "Net Realized", "Amount"].sum())
        canc = float(g.loc[g["financial_bucket"] == "Cancellation Leakage", "Amount"].sum())
        ret = float(g.loc[g["financial_bucket"] == "Return Leakage", "Amount"].sum())
        loss = float(g.loc[g["financial_bucket"] == "Loss Leakage", "Amount"].sum())
        rows.append({"Week": lbl, "Gross": round(gross, 2), "Net": round(net, 2),
                     "Cancelled": round(canc, 2), "Returned": round(ret, 2),
                     "Lost": round(loss, 2), "Leakage": round(canc + ret + loss, 2),
                     "Clearance%": round(net / gross * 100, 2) if gross else 0.0,
                     "Orders": len(g)})
    return pd.DataFrame(rows).sort_values("Week").reset_index(drop=True)


def main() -> None:
    """Load CSV, clean, aggregate, write single workbook."""
    print("Loading CSV ...", flush=True)
    df = pd.read_csv(RAW, low_memory=False)
    print(f"Rows: {len(df)}", flush=True)
    df = clean(df)
    weekly = weekly_summary(df)
    exc = df[df["financial_bucket"].isin(
        ["Cancellation Leakage", "Return Leakage", "Loss Leakage"])].copy()
    print(f"Weeks: {len(weekly)} | Exceptions: {len(exc)}", flush=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    print("Writing workbook ...", flush=True)
    with pd.ExcelWriter(OUT, engine="openpyxl") as w:
        ws0 = w.book.create_sheet("README", 0)
        ws0["A1"] = "Amazon Sales — ALL DATA in ONE Excel File"
        ws0["A1"].font = TITLE
        info = [
            f"Rows: {len(df)} | Weeks: {weekly['Week'].min()} to {weekly['Week'].max()}",
            f"Gross: Rs.{df['Amount'].sum():,.0f} | Net: Rs.{weekly['Net'].sum():,.0f}",
            "Tabs: 1) All Data  2) Weekly Summary  3) Leakage  4) Exceptions",
        ]
        for i, t in enumerate(info, start=3):
            ws0.cell(row=i, column=1, value=t)
        ws0.column_dimensions["A"].width = 90

        df.to_excel(w, sheet_name="All Data", index=False)
        weekly.to_excel(w, sheet_name="Weekly Summary", index=False)
        df.groupby(["financial_bucket", "Category"])["Amount"].sum().reset_index() \
            .to_excel(w, sheet_name="Leakage by Category", index=False)
        exc.to_excel(w, sheet_name="Exceptions", index=False)

        for name in ("All Data", "Weekly Summary", "Leakage by Category", "Exceptions"):
            ws = w.book[name]
            style_header(ws, 1, ws.max_column)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for c in range(1, min(ws.max_column + 1, 25)):
                ws.column_dimensions[get_column_letter(c)].width = 18
        ws2 = w.book["Weekly Summary"]
        ws2.conditional_formatting.add(
            f"H2:H{len(weekly) + 1}",
            CellIsRule(operator="lessThan", formula=["85"],
                       fill=PatternFill("solid", fgColor="E2E8F0")))
    print(f"DONE -> {OUT} ({OUT.stat().st_size // 1024} KB)", flush=True)


if __name__ == "__main__":
    main()

