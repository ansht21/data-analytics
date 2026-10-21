"""Tests for the 3-tab Excel workbook builder."""

from __future__ import annotations

import pandas as pd
from openpyxl import load_workbook
from src.reconciliation.engine import aggregate_weekly
from src.reconciliation.variance_analyzer import add_wow_variance
from src.reporting.excel_dashboard import build_workbook


def _weekly() -> pd.DataFrame:
    """Two-week synthetic KPI frame."""
    src = pd.DataFrame([
        ("W13", 2022, 13, "Net Realized", 1000.0),
        ("W13", 2022, 13, "Cancellation Leakage", 200.0),
        ("W14", 2022, 14, "Net Realized", 900.0),
        ("W14", 2022, 14, "Return Leakage", 100.0),
    ], columns=["week_label", "fiscal_year", "fiscal_week",
                "financial_bucket", "Amount"])
    return add_wow_variance(aggregate_weekly(src))


def test_workbook_tabs_and_content(settings) -> None:
    """Workbook writes 3 tabs with expected labels and data rows."""
    weekly = _weekly()
    exc = pd.DataFrame({"Order ID": ["X1"], "Status": ["Cancelled"],
                        "financial_bucket": ["Cancellation Leakage"],
                        "Amount": [200.0]})
    path = build_workbook(settings, weekly, exc, "W13-W14")
    assert path.exists() and path.stat().st_size > 4000
    wb = load_workbook(path)
    assert wb.sheetnames == ["Executive Dashboard", "Weekly Trends",
                             "Exception Audit Log"]
    trends = wb["Weekly Trends"]
    assert trends.cell(row=4, column=1).value == "W13"
    assert trends.cell(row=5, column=1).value == "W14"
    assert trends.cell(row=4, column=2).value == 1200.0  # gross W13
    audit = wb["Exception Audit Log"]
    assert audit.cell(row=5, column=1).value == "X1"
